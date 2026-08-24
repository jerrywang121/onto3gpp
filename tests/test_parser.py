"""Tests for the document parser."""

from pathlib import Path

from onto3gpp.parser import Section, parse_docx, parse_text


def test_parse_text_basic() -> None:
    text = """\
TS 23.501
5.6 PDU Session
5.6.1 General
The PDU session is established.
5.6.2 Procedures
The SMF selects a UPF.
"""
    sections = parse_text(text, spec="23.501", release="Rel-18")
    # Preamble section + one section per clause heading: 5.6, 5.6.1, 5.6.2
    assert len(sections) == 4
    assert sections[1].clause == "5.6"
    assert sections[2].clause == "5.6.1"
    assert sections[2].title == "General"
    assert "PDU session is established" in sections[2].text
    assert sections[2].spec == "23.501"
    assert sections[2].release == "Rel-18"


def test_parse_text_clause_hierarchy() -> None:
    text = """\
5.6 PDU Session
5.6.1 General
5.6.1.1 Sub
content here
"""
    sections = parse_text(text, spec="23.501", release="Rel-18")
    assert len(sections) == 3
    assert sections[2].clause == "5.6.1.1"
    assert sections[2].parent_clause == "5.6.1"


def test_parse_text_no_clauses() -> None:
    text = "Just a paragraph with no clause numbers."
    sections = parse_text(text, spec="23.501", release="Rel-18")
    assert len(sections) == 1
    assert sections[0].clause == ""
    assert sections[0].text == "Just a paragraph with no clause numbers."


def test_parse_docx(tmp_path: Path) -> None:
    from docx import Document

    doc = Document()
    doc.add_paragraph("5.6 PDU Session")
    doc.add_paragraph("5.6.1 General")
    doc.add_paragraph("The PDU session is established.")
    doc.add_paragraph("5.6.1.1\tSub\t114")  # TOC entry → dropped
    doc.add_paragraph("5.6.2 Procedures")
    doc.add_paragraph("The SMF selects a UPF.")
    path = tmp_path / "spec.docx"
    doc.save(str(path))

    sections = parse_docx(path, spec="38.300", release="Rel-18")
    # 5.6 + 5.6.1 + 5.6.2 (TOC entry dropped, no preamble)
    assert len(sections) == 3
    assert sections[0].clause == "5.6"
    assert sections[1].clause == "5.6.1"
    assert "PDU session is established" in sections[1].text
    assert sections[2].clause == "5.6.2"
